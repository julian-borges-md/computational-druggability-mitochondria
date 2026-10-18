#!/usr/bin/env python3
"""Render the main-text tables (1, 1b, 2, 3, 4) into R2_tables.docx from the
authoritative CSVs, mirroring the R1 tables document structure."""
import csv, os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
OUT=os.path.expanduser('~/r2/out'); TB=f'{OUT}/tables'
def rd(p): return list(csv.reader(open(p)))

doc=Document()
st=doc.styles['Normal']; st.font.name='Calibri'; st.font.size=Pt(10)

def H(text, size=13, space=6):
    p=doc.add_paragraph(); r=p.add_run(text); r.bold=True; r.font.size=Pt(size)
    p.space_after=Pt(space); return p
def note(text):
    p=doc.add_paragraph(); r=p.add_run(text); r.italic=True; r.font.size=Pt(8.5)
    r.font.color.rgb=RGBColor(0x44,0x44,0x44); return p

def table(rows, header=True, fontsize=8.5):
    t=doc.add_table(rows=len(rows), cols=len(rows[0])); t.style='Light Grid Accent 1'
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,row in enumerate(rows):
        for j,val in enumerate(row):
            c=t.cell(i,j); c.text=''
            para=c.paragraphs[0]; run=para.add_run(str(val)); run.font.size=Pt(fontsize)
            if header and i==0: run.bold=True
    return t

# ---------------- Title ----------------
title=doc.add_paragraph(); r=title.add_run('Revised Manuscript Tables, revision R2'); r.bold=True; r.font.size=Pt(15)
doc.add_paragraph('Manuscript CSBR-D-26-00031, Computational and Structural Biotechnology Reports').runs[0].font.size=Pt(10)
doc.add_paragraph('Julian Yin Vieira Borges, MD. 30 September 2026').runs[0].font.size=Pt(10)
note('Every value is regenerated from the deposited data by the reproduction scripts. The structural '
     'dimension is scored at the a priori therapeutic site (Supplementary Table S1) in two specifications, '
     'the isolated chain and the biological assembly; D3 is the curated pChEMBL compound count.')

# ---------------- Table 1 ----------------
H('Table 1. Five dimension assessment for all 25 targets, site matched')
note('D1 is the fpocket Druggability Score of the pocket overlapping the a priori therapeutic site, '
     'reported for the isolated chain and the biological assembly. D2 is the site pocket volume. '
     'D3 is the number of distinct ChEMBL compounds with pChEMBL >= 5. D4 and D5 are ordinal (0 to 3). '
     'NA marks the eight targets with no defined small molecule site. Composite is the weight renormalised '
     'mean over available dimensions.')
table(rd(f'{TB}/Table1_five_dimension.csv'))
note('D4: 3 approved drug at the assessed site, 2 clinical stage, 1 validated tool or preclinical, 0 none. '
     'D5: 3 no targeting required, 2 matrix targeting, 1 IMS or membrane embedded, 0 gene therapy primary or no viable route.')
doc.add_paragraph()

# ---------------- Table 1b provenance ----------------
H('Table 1b. Structure provenance for all 25 targets')
prov=rd(f'{OUT}/provenance.csv')
hdr=['Target','Protein','PDB / model','Res A','Organism','Mutation','Coverage %','Chains kept / entity','Het removed']
table([hdr]+prov[1:])
note('Coverage is residues resolved in the best retained chain divided by the UniProt reference length. '
     'Four assessments fall below 60 per cent: LONP1 19.8, OPA1 34.4, DNM1L 46.5, KEAP1 46.6. '
     'Coverage is uncorrelated with the pocket scores across the panel. AlphaFold models carry no resolution.')
doc.add_paragraph()

# ---------------- Table 2 delivery ----------------
H('Table 2. Delivery profile partition')
deliv=[['Delivery category','n','Strategy','Targets'],
 ['No mitochondrial targeting required','5','Standard systemic delivery','DNM1L, PINK1, NFE2L2 via KEAP1, VDAC1, PPARGC1A'],
 ['Matrix targeting','9','Matrix targeting peptide or TPP cation conjugation','NDUFV1, SDHA, ATP5F1A, POLG, SOD2, LONP1, CLPP, TWNK, TFAM'],
 ['Intermembrane space or membrane embedded','7','Matrix targeting peptide or cell penetrating peptide','MT-CYB, MT-CO1, UQCRC2, COX4I1, OPA1, SLC25A4, TOMM40'],
 ['Gene therapy primary','3','Allotopic expression, heteroplasmy shifting, base editing','MT-ND1, MT-ND4, MT-ATP6'],
 ['No viable small molecule strategy','1','Assembly factor or chaperone approach','TIMM23'],
 ['Total','25','','All 25 targets assigned exactly once']]
table(deliv)
doc.add_paragraph()

# ---------------- Table 3 tier-1 ----------------
H('Table 3. All 15 Tier 1 targets ranked, with the two stage selection basis')
note('Stage 1 ranks by composite score (overall rank shown). Stage 2 fails a target whose only chemical '
     'precedent is inhibition where the disease mechanism requires rescue, and records the disease and '
     'predominant variant mechanism for the first wave. D1 is the site matched assembly score.')
table(rd(f'{TB}/Table3_tier1_ranked.csv'), fontsize=8)
note('Revised first wave: DNM1L, NFE2L2 via KEAP1, LONP1, SDHA, CLPP. Stable across the isolated and '
     'assembly structural specifications and in more than 94 per cent of randomised dimension weightings for four of the five.')
doc.add_paragraph()

# ---------------- Table 4 tier encoding ----------------
H('Table 4. What the tier label encodes')
note('Spearman rank correlations. Permutation p (200,000 permutations, seed 20260930) is the primary '
     'inferential value; percentile bootstrap 95 per cent intervals (100,000 resamples) are shown; '
     'Benjamini-Hochberg q controls the seven test family. The first five rows are dimension against tier; '
     'the last two are the central independence tests.')
table(rd(f'{TB}/Table4_tier_encoding.csv'), fontsize=8)
note('At n = 25 the minimum detectable correlation at 80 per cent power is 0.535, so every null is a null '
     'at this panel size. Structure to precedent at the a priori site: +0.117 isolated chain, -0.100 biological '
     'assembly, +0.052 globally best pocket originally reported.')

doc.save(os.path.expanduser('~/r2/build/R2_tables.docx'))
print("wrote build/R2_tables.docx")
